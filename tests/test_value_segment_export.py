"""Endpoint tests for the value-segments downloads on the Dashboard:
/api/value-segments/export.xlsx (the whole card — a Summary sheet plus one
sheet per segment) and /api/value-segments/export.csv (one tile's SKU list).
"""

import csv
import io
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.gowri_proj import db
from src.gowri_proj.analysis import SEGMENT_ORDER
from src.gowri_proj.parser import TIDY_COLUMNS, ReportMeta
from src.gowri_proj.webapp import create_app


def _meta(start, end):
    return ReportMeta(
        company="TEST PHARMACY", location="TEST CITY",
        period_start=date.fromisoformat(start), period_end=date.fromisoformat(end),
    )


def _row(sku, closing, value, sales, brand="BRAND"):
    return {
        "brand": brand, "sku": sku, "opening_stock": closing + sales, "purchase": 0.0,
        "purchase_free": 0.0, "other_receipt": 0.0, "sales": sales, "sales_free": 0.0,
        "other_issue": 0.0, "closing_stock": closing, "value": value,
    }


@pytest.fixture
def client(tmp_path):
    app = create_app(db_path=str(tmp_path / "test.db"), uploads_dir=str(tmp_path / "uploads"))
    with app.test_client() as c:
        yield c, app


def _seed(app):
    # One 30-day report. BIG dominates value (tier A); the rest fall to B/C.
    # BIG and MID sell a healthy amount (fast); SLOW has stock but no sales at
    # all (infinite days of cover -> overstock -> "slow"); GONE is out of
    # stock and must not appear in any segment.
    rows = [
        _row("BIG", closing=10, value=1000.0, sales=10),
        _row("MID", closing=10, value=200.0, sales=10),
        _row("SLOW", closing=5, value=50.0, sales=0),
        _row("GONE", closing=0, value=0.0, sales=5),
    ]
    with db.connect(app.config["DB_PATH"]) as conn:
        db.import_report(
            conn, pd.DataFrame(rows, columns=TIDY_COLUMNS), _meta("2026-06-01", "2026-06-30"),
            "jun.xls",
        )


def test_xlsx_has_summary_plus_one_sheet_per_segment(client):
    c, app = client
    _seed(app)
    resp = c.get("/api/value-segments/export.xlsx")
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'filename="value_segments_2026-06-30.xlsx"' in resp.headers["Content-Disposition"]

    wb = load_workbook(io.BytesIO(resp.data))
    assert len(wb.sheetnames) == 1 + len(SEGMENT_ORDER)
    assert wb.sheetnames[0] == "Summary"


def test_xlsx_summary_counts_match_segment_sheets(client):
    c, app = client
    _seed(app)
    wb = load_workbook(io.BytesIO(c.get("/api/value-segments/export.xlsx").data))

    summary = list(wb["Summary"].iter_rows(values_only=True))
    header_idx = next(i for i, r in enumerate(summary) if r and r[0] == "Tier")
    header = summary[header_idx]
    total = 0
    for r in summary[header_idx + 1:]:
        rec = dict(zip(header, r))
        sheet_rows = list(wb[rec["Sheet"]].iter_rows(values_only=True))[1:]
        assert len(sheet_rows) == rec["SKUs"]
        assert rec["Recommendation"]
        total += rec["SKUs"]
    # Every stocked SKU segmented exactly once; GONE (out of stock) excluded.
    assert total == 3


def test_xlsx_segment_sheet_columns_and_infinite_cover_is_blank(client):
    c, app = client
    _seed(app)
    wb = load_workbook(io.BytesIO(c.get("/api/value-segments/export.xlsx").data))
    found = {}
    for name in wb.sheetnames[1:]:
        rows = list(wb[name].iter_rows(values_only=True))
        header = rows[0]
        assert header[:2] == ("Brand", "SKU")
        assert "Days Since Activity" in header
        for r in rows[1:]:
            found[r[1]] = dict(zip(header, r))
    assert set(found) == {"BIG", "MID", "SLOW"}
    assert found["SLOW"]["Days of Cover"] is None  # no demand: blank, not "inf"
    assert found["BIG"]["Closing Stock (now)"] == 10


def test_csv_returns_one_segments_rows(client):
    c, app = client
    _seed(app)
    wb = load_workbook(io.BytesIO(c.get("/api/value-segments/export.xlsx").data))
    # Find which tier SLOW landed in, then ask for exactly that tile.
    tier = next(
        n.split(" - ")[0] for n in wb.sheetnames[1:]
        if any(r[1] == "SLOW" for r in wb[n].iter_rows(min_row=2, values_only=True))
    )
    resp = c.get(f"/api/value-segments/export.csv?tier={tier}&movement=slow")
    assert resp.status_code == 200
    assert resp.mimetype == "text/csv"
    rows = list(csv.reader(io.StringIO(resp.data.decode("utf-8-sig"))))
    assert rows[0][:2] == ["Brand", "SKU"]
    assert [r[1] for r in rows[1:]] == ["SLOW"]
    days_cover_col = rows[0].index("Days of Cover")
    assert rows[1][days_cover_col] == ""


def test_csv_rejects_unknown_segment(client):
    c, app = client
    _seed(app)
    assert c.get("/api/value-segments/export.csv?tier=Z&movement=slow").status_code == 400
    assert c.get("/api/value-segments/export.csv?tier=A&movement=bogus").status_code == 400


def test_exports_404_with_no_data(client):
    c, _ = client
    assert c.get("/api/value-segments/export.xlsx").status_code == 404
    assert c.get("/api/value-segments/export.csv?tier=A&movement=fast").status_code == 404


def test_dashboard_page_links_to_the_downloads(client):
    c, app = client
    _seed(app)
    html = c.get("/dashboard").data.decode()
    assert "/api/value-segments/export.xlsx" in html
    assert "/api/value-segments/export.csv" in html
