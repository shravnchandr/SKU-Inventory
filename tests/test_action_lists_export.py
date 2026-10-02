"""The Dashboard's "Download all lists (Excel)": every action list in full
(not the 300 rows shown on screen), one sheet each, plus a Summary — the
same workbook the CLI writes."""

import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.gowri_proj import db
from src.gowri_proj.parser import TIDY_COLUMNS, ReportMeta
from src.gowri_proj.webapp import create_app


@pytest.fixture
def client(tmp_path):
    app = create_app(db_path=str(tmp_path / "t.db"), uploads_dir=str(tmp_path / "u"))
    with app.test_client() as c:
        yield c, app


def test_404_with_no_data(client):
    c, _ = client
    assert c.get("/api/action-lists/export.xlsx").status_code == 404


def test_every_row_of_every_list(client):
    c, app = client
    # 400 out-of-stock SKUs — more than the 300 the page shows at once.
    rows = [
        {"brand": "B", "sku": f"OOS {i:03}", "opening_stock": 5.0, "purchase": 0.0, "purchase_free": 0.0,
         "other_receipt": 0.0, "sales": 5.0, "sales_free": 0.0, "other_issue": 0.0,
         "closing_stock": 0.0, "value": 0.0}
        for i in range(400)
    ]  # fmt: skip
    with db.connect(app.config["DB_PATH"]) as conn:
        db.import_report(
            conn, pd.DataFrame(rows, columns=TIDY_COLUMNS),
            ReportMeta("T", "T", pd.Timestamp("2026-07-01").date(), pd.Timestamp("2026-07-31").date()),
            "jul.xls",
        )  # fmt: skip
    resp = c.get("/api/action-lists/export.xlsx")
    assert resp.status_code == 200
    assert 'filename="action_lists_2026-07-31.xlsx"' in resp.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["Summary", "Out of Stock", "Low Stock", "Dead Stock", "Overstock"]
    assert wb["Out of Stock"].max_row - 1 == 400
    assert "/api/action-lists/export.xlsx" in c.get("/dashboard").data.decode()
