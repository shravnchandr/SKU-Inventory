"""Both parse_stock_statement and parse_item_list trust column positions by
fixed offset (COL_VALUE, IL_COL_LONG_NAME, etc.) once the header row is
found. Uploading the wrong kind of spreadsheet — one with too few columns
to ever reach those offsets — used to fail with a bare ``KeyError`` the
moment the row loop hit a missing column: technically accurate, meaningless
to whoever's looking at it. Both parsers now check the column count up
front and raise a plain-English ValueError instead.

Also covers the sibling failure mode: a file pandas/xlrd/openpyxl can't even
open at all (wrong file type entirely, or corrupted/partially
saved/downloaded) — previously surfaced whatever raw engine exception the
read failed with (CompDocError, BadZipFile, a bare "you must specify an
engine manually" ValueError); _read_excel_safely now turns any of those into
one plain message too.
"""

import openpyxl
import pytest

from src.gowri_proj.parser import parse_item_list, parse_stock_statement


def test_stock_statement_with_too_few_columns_gets_a_clear_error(tmp_path):
    path = tmp_path / "wrong_export.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["TEST PHARMACY"])
    ws.append(["TEST CITY"])
    ws.append(["Stock Statement from 01/Jun/2026 to 30/Jun/2026"])
    # Only 5 columns — nowhere near the 15 a real stock statement export has
    # (this is the shape you'd get from, say, a differently-formatted report
    # saved from the same POS system).
    ws.append(["Item", "Qty", "Rate", "Amount", "Tax"])
    ws.append(["SOME SKU", 10, 5, 50, 2])
    wb.save(path)

    with pytest.raises(ValueError, match="only has 5 column"):
        parse_stock_statement(str(path))


def test_item_list_with_too_few_columns_gets_a_clear_error(tmp_path):
    path = tmp_path / "wrong_export.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    sheet2 = wb.create_sheet("Sheet2")
    sheet2.append(["TEST PHARMACY"])
    sheet2.append(["TEST CITY"])
    sheet2.append(["Item List as on 09/08/2026"])
    sheet2.append(["Code", "Product", "Packing"])
    sheet2.append(["C001", "SOME ITEM", "10S"])
    wb.save(path)

    with pytest.raises(ValueError, match="only has 3 column"):
        parse_item_list(str(path))


def test_stock_statement_from_a_non_excel_file_gets_a_clear_error(tmp_path):
    # e.g. a web page or PDF saved/renamed with an .xls extension by mistake.
    path = tmp_path / "not_actually_excel.xls"
    path.write_text("<html>this is not a spreadsheet</html>")

    with pytest.raises(ValueError, match="Couldn't open this file as an Excel spreadsheet"):
        parse_stock_statement(str(path))


def test_item_list_from_a_non_excel_file_gets_a_clear_error(tmp_path):
    path = tmp_path / "not_actually_excel.xlsx"
    path.write_text("this isn't a real xlsx either")

    with pytest.raises(ValueError, match="Couldn't open this file as an Excel spreadsheet"):
        parse_item_list(str(path))


def test_stock_statement_from_a_truncated_file_gets_a_clear_error(tmp_path):
    # A download/save that got cut off partway through — the file is a
    # genuine (partial) .xls, not garbage, but still unreadable. Built by
    # truncating a real fixture rather than hand-crafting bytes, since the
    # exact corrupted-OLE2-container shape is what xlrd actually chokes on.
    good_path = tmp_path / "good.xls"
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet1")
    for r in range(200):
        for c in range(15):
            ws.write(r, c, f"padding {r}-{c}")
    wb.save(str(good_path))

    truncated_path = tmp_path / "truncated.xls"
    data = good_path.read_bytes()
    truncated_path.write_bytes(data[: len(data) // 2])

    with pytest.raises(ValueError, match="Couldn't open this file as an Excel spreadsheet"):
        parse_stock_statement(str(truncated_path))


def test_item_list_row_without_a_code_is_skipped_not_keyed_as_nan(tmp_path):
    # An item can only be matched by its code. A row with a name but no code
    # used to get the code "nan" — two of them collided on the primary key
    # and blocked the whole item-list upload.
    from src.gowri_proj.parser import parse_item_list
    from tests.test_webapp_catalog_upload import _write_item_list

    path = tmp_path / "il.xlsx"
    _write_item_list(path, [
        ("brand", "BRAND A"),
        ("item", "X1", "X TAB", "10", 1.0, 1.0, 12.0, "3004", None),
        ("item", None, "NO CODE ONE", "10", 1.0, 1.0, 12.0, "3004", None),
        ("item", None, "NO CODE TWO", "10", 1.0, 1.0, 12.0, "3004", None),
    ])  # fmt: skip
    df, _ = parse_item_list(str(path))
    assert list(df["code"]) == ["X1"]
    assert list(df["brand"]) == ["BRAND A"]
